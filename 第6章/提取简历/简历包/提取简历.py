# -*- coding: utf-8 -*-
from docx import Document
import pandas as pd
import glob
files=glob.glob(r'G:\示例\第6章\提取简历\简历包\docx\*.docx')
data=[]
for file in files:
    doc = Document(file)
    table = doc.tables[0]
    value=[table.cell(0,1).text,table.cell(0,4).text,
    table.cell(0,6).text,table.cell(1,1).text,table.cell(1,4).text,
    table.cell(1,6).text,table.cell(2,1).text,table.cell(2,4).text,
    table.cell(2,6).text,table.cell(3,1).text,table.cell(3,4).text,
    table.cell(3,6).text,table.cell(4,2).text,table.cell(4,6).text,
    table.cell(5,2).text,table.cell(5,6).text,table.cell(6,2).text,
    table.cell(7,2).text,table.cell(8,2).text,table.cell(9,2).text,
    table.cell(10,2).text,table.cell(11,2).text,table.cell(12,2).text]
    data.append(value)
    for rel in doc.part._rels:
        rel = doc.part._rels[rel]            
        if "image" not in rel.target_ref:
            continue
        imgName =r'G:\示例\第6章\提取简历\照片集合'+"\\" +table.cell(0,1).text+".png"
        with open(imgName,"wb") as f:
            f.write(rel.target_part.blob)

name=["姓名", "性别", "出生年月", "籍贯", "民族", "政治面貌", "学历", "学制", "外语水平", "身高", "体重", "计算机水平", "毕业学校", "专业", "E-mail", "手机", "社会实践经历", "在校期间获奖情况", "担任主要社会工作", "外语能力", "计算机能力", "兴趣爱好", "自我评价"]
df = pd.DataFrame(data, columns=name)
df.to_excel(r"G:\示例\第6章\提取简历\提取简历结果.xlsx", index=False)

    
    
    
    
    
